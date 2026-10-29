# -*- coding: utf-8 -*-
"""
=====================================================================
Word Report — سند Word (docx) با قالب‌بندی حرفه‌ای و راست‌به‌چپ
=====================================================================
شامل جلد رنگی، شاخص‌های کلیدی، جدول‌های راه‌راه، تنظیمات RTL
پاراگراف‌ها و جدول‌ها با استفاده از python-docx.
=====================================================================
"""
from __future__ import annotations

import io
from typing import Any, Dict, List

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from .theme import PALETTE, fa_num

p = PALETTE


def _rgb(hex_color: str) -> RGBColor:
    hex_color = hex_color.lstrip("#")
    return RGBColor(
        int(hex_color[0:2], 16),
        int(hex_color[2:4], 16),
        int(hex_color[4:6], 16),
    )


def _set_rtl(paragraph) -> None:
    """فعال‌سازی راست‌به‌چپ برای پاراگراف."""
    pPr = paragraph._p.get_or_add_pPr()
    bidi = pPr.find(qn("w:bidi"))
    if bidi is None:
        bidi = OxmlElement("w:bidi")
        pPr.append(bidi)
    bidi.set(qn("w:val"), "1")


def _set_run_font(run, name: str = "Vazirmatn", size: int = 11,
                  color: str | None = None, bold: bool = False) -> None:
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = _rgb(color)
    # فونت برای متن فارسی (complex script)
    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:cs"), name)
    run.font.cs_bold = bold


def _cell_bg(cell, hex_color: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_color.lstrip("#"))
    tcPr.append(shd)


def _add_para(doc, text: str, size: int = 11, color: str = p["ink"],
              bold: bool = False, align=WD_ALIGN_PARAGRAPH.RIGHT,
              space_after: int = 6) -> Any:
    para = doc.add_paragraph()
    para.alignment = align
    para.paragraph_format.space_after = Pt(space_after)
    _set_rtl(para)
    run = para.add_run(text)
    _set_run_font(run, size=size, color=color, bold=bold)
    return para


def _add_heading(doc, text: str, color: str) -> None:
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    para.paragraph_format.space_before = Pt(14)
    para.paragraph_format.space_after = Pt(6)
    _set_rtl(para)
    run = para.add_run("■ " + text)
    _set_run_font(run, size=14, color=color, bold=True)
    # خط زیرین
    pPr = para._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "2")
    bottom.set(qn("w:color"), p["line"].lstrip("#"))
    pBdr.append(bottom)
    pPr.append(pBdr)


def _add_table(doc, headers: List[str], rows: List[List[str]],
               header_bg: str, widths: List[int] | None = None) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"

    # header
    for col, h in enumerate(headers):
        cell = table.cell(0, col)
        _cell_bg(cell, header_bg)
        para = cell.paragraphs[0]
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _set_rtl(para)
        run = para.add_run(h)
        _set_run_font(run, size=10, color=p["primary_dark"], bold=True)

    # body
    for r, row in enumerate(rows, start=1):
        for col, value in enumerate(row):
            cell = table.cell(r, col)
            if r % 2 == 0:
                _cell_bg(cell, p["bg"])
            para = cell.paragraphs[0]
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER if col > 0 \
                else WD_ALIGN_PARAGRAPH.RIGHT
            _set_rtl(para)
            run = para.add_run(str(value))
            _set_run_font(
                run, size=9.5, color=p["ink"], bold=(col == 0)
            )

    if widths:
        for col, w in enumerate(widths):
            for row_cells in table.rows:
                row_cells.cells[col].width = Pt(w)


def build_docx(data: Dict[str, Any]) -> tuple[bytes, str, str]:
    kpis: Dict[str, int] = data["kpis"]
    courses: List[Dict[str, Any]] = data["courses"]
    quizzes: List[Dict[str, Any]] = data["quizzes"]
    scheduler: Dict[str, Any] | None = data.get("scheduler")
    scheduler_courses: List[Dict[str, Any]] = data.get("scheduler_courses") or []

    doc = Document()

    # فونت پیش‌فرض سند
    style = doc.styles["Normal"]
    style.font.name = "Vazirmatn"
    style.font.size = Pt(11)
    style.element.rPr.rFonts.set(qn("w:cs"), "Vazirmatn")

    # ==============================================================
    # جلد
    # ==============================================================
    cover_table = doc.add_table(rows=1, cols=1)
    cover_cell = cover_table.cell(0, 0)
    _cell_bg(cover_cell, p["primary_dark"])

    para = cover_cell.paragraphs[0]
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_before = Pt(30)
    _set_rtl(para)
    run = para.add_run(data["title"])
    _set_run_font(run, size=24, color="#FFFFFF", bold=True)

    para2 = cover_cell.add_paragraph()
    para2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_rtl(para2)
    run2 = para2.add_run(data["subtitle"])
    _set_run_font(run2, size=12, color="#E0E7FF")

    para3 = cover_cell.add_paragraph()
    para3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para3.paragraph_format.space_after = Pt(30)
    _set_rtl(para3)
    run3 = para3.add_run(f"🗓 {data['date_fa']}   ·   🕒 {data['generated_at']}")
    _set_run_font(run3, size=10, color="#C7D2FE")

    doc.add_paragraph()

    # ==============================================================
    # شاخص‌های کلیدی
    # ==============================================================
    _add_heading(doc, "شاخص‌های کلیدی سامانه", p["primary"])
    _add_table(
        doc,
        ["درس", "سرفصل", "دقیقه آموزش", "محتوای آموزشی", "آزمون", "سؤال"],
        [[
            fa_num(kpis["courses"]),
            fa_num(kpis["topics"]),
            fa_num(kpis["minutes"]),
            fa_num(kpis["contents"]),
            fa_num(kpis["quizzes"]),
            fa_num(kpis["questions"]),
        ]],
        header_bg=p["primary_light"].lstrip("#"),
    )
    doc.add_paragraph()

    # ==============================================================
    # ترکیب سؤال‌ها
    # ==============================================================
    split = data["question_split"]
    total_q = sum(i["value"] for i in split)
    _add_heading(doc, "ترکیب سؤال‌ها بر اساس نوع", p["cyan"])
    if total_q:
        lines = [
            f"• {i['label']}: {fa_num(i['value'])} سؤال "
            f"({fa_num(round(i['value'] / total_q * 100))}٪)"
            for i in split
        ]
        for line in lines:
            _add_para(doc, line, size=10.5)
    else:
        _add_para(doc, "هنوز سؤالی ثبت نشده است.", size=10.5,
                  color=p["muted"])

    # ==============================================================
    # فهرست دروس
    # ==============================================================
    _add_heading(doc, "فهرست دروس و شاخص‌ها", p["amber"])
    if courses:
        rows = [
            [
                c["title"][:30],
                fa_num(c["topic_count"]),
                fa_num(c["minutes"]),
                fa_num(c["content_count"]),
                fa_num(c["quiz_count"]),
                fa_num(c["question_count"]),
                "فعال" if c["is_active"] else "غیرفعال",
            ]
            for c in courses[:20]
        ]
        _add_table(
            doc,
            ["عنوان درس", "سرفصل", "دقیقه", "محتوا", "آزمون", "سؤال", "وضعیت"],
            rows,
            header_bg=p["primary_light"].lstrip("#"),
            widths=[110, 40, 40, 45, 40, 40, 50],
        )
    else:
        _add_para(doc, "هنوز درسی ثبت نشده است.", size=10.5, color=p["muted"])
    doc.add_paragraph()

    # ==============================================================
    # آزمون‌ها
    # ==============================================================
    _add_heading(doc, "آزمون‌های سامانه", p["accent"])
    if quizzes:
        rows = [
            [q["title"][:28], q["course"][:20], fa_num(q["question_count"])]
            for q in quizzes[:12]
        ]
        _add_table(
            doc,
            ["عنوان آزمون", "درس", "تعداد سؤال"],
            rows,
            header_bg=p["primary_light"].lstrip("#"),
        )
    else:
        _add_para(doc, "هنوز آزمونی ساخته نشده است.", size=10.5,
                  color=p["muted"])
    doc.add_paragraph()

    # ==============================================================
    # یکپارچگی سرویس‌ها
    # ==============================================================
    if scheduler:
        _add_heading(doc, "یکپارچگی سرویس‌ها (Service Integration)", p["teal"])
        if scheduler.get("available"):
            _add_para(
                doc,
                f"وضعیت اتصال به سامانه برنامه‌ریزی درسی: متصل ✓ — "
                f"{fa_num(scheduler.get('course_count', 0))} درس از سرویس خوانده شد.",
                size=10.5, color="#15803D", bold=True,
            )
        else:
            _add_para(
                doc,
                "در زمان تولید گزارش، سرویس برنامه‌ریزی درسی پاسخ نداد؛ "
                "گزارش بر پایهٔ داده‌های محلی سامانه ساخته شده است.",
                size=10.5, color="#B91C1C", bold=True,
            )
        _add_para(doc, f"آدرس سرویس: {scheduler.get('base_url', '—')}",
                  size=9.5, color=p["muted"])

        if scheduler_courses:
            rows = [
                [c["title"][:34], str(c.get("code") or "—"),
                 fa_num(c.get("credits", "—"))]
                for c in scheduler_courses[:10]
            ]
            _add_table(
                doc,
                ["درس (از سرویس scheduler)", "کد", "واحد"],
                rows,
                header_bg="CCFBF1",
            )

    doc.add_paragraph()
    _add_para(
        doc,
        f"این گزارش به‌صورت خودکار توسط «سامانه هوشمند آموزش» تولید شده است — {data['date_fa']}",
        size=9, color=p["muted"], align=WD_ALIGN_PARAGRAPH.CENTER,
    )

    buf = io.BytesIO()
    doc.save(buf)
    filename = (
        f"smart-education-report-"
        f"{data['generated_at'].replace('-', '').replace(':', '').replace(' ', '-')}.docx"
    )
    return (
        buf.getvalue(),
        filename,
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
