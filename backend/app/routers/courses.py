import csv
import io
from pathlib import Path
from typing import Any, Dict, List, Tuple

from fastapi import (
    APIRouter,
    Depends,
    File,
    Form,
    HTTPException,
    UploadFile,
    status,
)
from openpyxl import load_workbook
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from ..database import get_db
from ..models import Course, CourseTopic
from ..schemas import CourseCreate, CourseResponse, TopicResponse


# ============================================================
# Router
# ============================================================

router = APIRouter(
    prefix="/courses",
    tags=["courses"],
)


# ============================================================
# تنظیمات
# ============================================================

MAX_FILE_SIZE = 10 * 1024 * 1024

ALLOWED_EXTENSIONS = {
    ".xlsx",
    ".csv",
    ".tsv",
    ".txt",
}

REQUIRED_COLUMNS = {
    "id",
    "course_id",
    "title",
    "description",
    "session_minute",
    "sort_order",
    "is_active",
}

PERSIAN_DIGITS = "۰۱۲۳۴۵۶۷۸۹"
ARABIC_DIGITS = "٠١٢٣٤٥٦٧٨٩"
ENGLISH_DIGITS = "0123456789"


# ============================================================
# خطا
# ============================================================

def raise_bad_request(message: str) -> None:
    raise HTTPException(
        status_code=status.HTTP_400_BAD_REQUEST,
        detail=message,
    )


# ============================================================
# پاک‌سازی و تبدیل داده
# ============================================================

def normalize_digits(value: Any) -> Any:
    if not isinstance(value, str):
        return value

    table = str.maketrans(
        PERSIAN_DIGITS + ARABIC_DIGITS,
        ENGLISH_DIGITS + ENGLISH_DIGITS,
    )

    return value.translate(table)


def clean_header(value: Any) -> str:
    if value is None:
        return ""

    return (
        str(value)
        .replace("\ufeff", "")
        .replace("\r", " ")
        .replace("\n", " ")
        .strip()
        .lower()
    )


def clean_value(value: Any) -> Any:
    if value is None:
        return ""

    if isinstance(value, str):
        value = (
            value
            .replace("\ufeff", "")
            .replace("\r", " ")
            .replace("\n", " ")
            .strip()
        )

        return normalize_digits(value)

    return value


def is_empty_value(value: Any) -> bool:
    if value is None:
        return True

    return str(value).strip() == ""


def convert_to_int(
    value: Any,
    field_name: str,
    row_number: int,
    default: int | None = None,
    minimum: int | None = None,
) -> int:
    value = clean_value(value)

    if is_empty_value(value):
        if default is not None:
            return default

        raise_bad_request(
            f"ستون «{field_name}» در سطر "
            f"{row_number} الزامی است."
        )

    try:
        number = float(value)

        if not number.is_integer():
            raise ValueError

        result = int(number)

    except (TypeError, ValueError):
        raise_bad_request(
            f"مقدار ستون «{field_name}» در سطر "
            f"{row_number} باید عدد صحیح باشد: {value}"
        )

    if minimum is not None and result < minimum:
        raise_bad_request(
            f"مقدار ستون «{field_name}» در سطر "
            f"{row_number} باید حداقل {minimum} باشد."
        )

    return result


def convert_to_bool(
    value: Any,
    row_number: int,
    default: bool = True,
) -> bool:
    if value is None or is_empty_value(value):
        return default

    if isinstance(value, bool):
        return value

    text = str(normalize_digits(value)).strip().lower()

    true_values = {
        "1",
        "true",
        "yes",
        "y",
        "on",
        "active",
        "فعال",
        "بله",
    }

    false_values = {
        "0",
        "false",
        "no",
        "n",
        "off",
        "inactive",
        "غیرفعال",
        "خیر",
    }

    if text in true_values:
        return True

    if text in false_values:
        return False

    raise_bad_request(
        f"مقدار ستون «is_active» در سطر "
        f"{row_number} معتبر نیست: {value}"
    )

    return default


# ============================================================
# اعتبارسنجی فایل
# ============================================================

def validate_columns(columns: List[Any]) -> List[str]:
    headers = [
        clean_header(column)
        for column in columns
    ]

    if not headers:
        raise_bad_request("فایل هیچ ستونی ندارد.")

    if any(not header for header in headers):
        raise_bad_request(
            "یکی از ستون‌های فایل نام ندارد."
        )

    duplicated_columns = {
        header
        for header in headers
        if headers.count(header) > 1
    }

    if duplicated_columns:
        raise_bad_request(
            "نام ستون‌های تکراری در فایل: "
            f"{sorted(duplicated_columns)}"
        )

    actual_columns = set(headers)
    missing_columns = REQUIRED_COLUMNS - actual_columns

    if missing_columns:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail={
                "message": (
                    "برخی ستون‌های ضروری در فایل "
                    "وجود ندارند."
                ),
                "missing_columns": sorted(missing_columns),
                "received_columns": sorted(actual_columns),
                "required_columns": sorted(REQUIRED_COLUMNS),
            },
        )

    return headers


def make_record(
    headers: List[str],
    values: List[Any],
    row_number: int,
) -> Dict[str, Any]:
    normalized_values = list(values)

    if len(normalized_values) < len(headers):
        normalized_values.extend(
            [""] * (len(headers) - len(normalized_values))
        )

    if len(normalized_values) > len(headers):
        normalized_values = normalized_values[:len(headers)]

    record = {
        header: clean_value(normalized_values[index])
        for index, header in enumerate(headers)
    }

    if all(
        is_empty_value(value)
        for value in record.values()
    ):
        return {}

    record["_row_number"] = row_number

    return record


# ============================================================
# خواندن Excel
# ============================================================

def read_excel_file(
    file_bytes: bytes,
) -> List[Dict[str, Any]]:
    try:
        workbook = load_workbook(
            filename=io.BytesIO(file_bytes),
            read_only=True,
            data_only=True,
        )

    except Exception as exc:
        raise_bad_request(
            "فایل Excel قابل خواندن نیست. "
            f"جزئیات: {exc}"
        )

    try:
        worksheet = workbook.active
        rows = worksheet.iter_rows(values_only=True)

        try:
            header_row = next(rows)

        except StopIteration:
            raise_bad_request(
                "فایل Excel خالی است."
            )

        headers = validate_columns(
            list(header_row)
        )

        records: List[Dict[str, Any]] = []

        for row_number, row_values in enumerate(
            rows,
            start=2,
        ):
            record = make_record(
                headers=headers,
                values=list(row_values),
                row_number=row_number,
            )

            if record:
                records.append(record)

        return records

    finally:
        workbook.close()


# ============================================================
# خواندن فایل متنی
# ============================================================

def decode_text_file(file_bytes: bytes) -> str:
    encodings = (
        "utf-8-sig",
        "utf-8",
        "cp1256",
        "windows-1256",
        "cp1252",
        "latin-1",
    )

    for encoding in encodings:
        try:
            return file_bytes.decode(encoding)

        except UnicodeDecodeError:
            continue

    raise_bad_request(
        "Encoding فایل قابل تشخیص نیست. "
        "فایل را با UTF-8 ذخیره کنید."
    )

    return ""


def detect_delimiter(
    text: str,
    extension: str,
) -> str:
    if extension == ".tsv":
        return "\t"

    sample = text[:8192]

    try:
        dialect = csv.Sniffer().sniff(
            sample,
            delimiters=",;\t|",
        )

        return dialect.delimiter

    except csv.Error:
        if "\t" in sample:
            return "\t"

        if ";" in sample:
            return ";"

        if "|" in sample:
            return "|"

        return ","


def read_text_file(
    file_bytes: bytes,
    extension: str,
) -> List[Dict[str, Any]]:
    text = decode_text_file(file_bytes)

    if not text.strip():
        raise_bad_request(
            "فایل متنی خالی است."
        )

    delimiter = detect_delimiter(
        text=text,
        extension=extension,
    )

    reader = csv.reader(
        io.StringIO(text),
        delimiter=delimiter,
    )

    try:
        header_row = next(reader)

    except StopIteration:
        raise_bad_request(
            "فایل متنی خالی است."
        )

    headers = validate_columns(header_row)
    records: List[Dict[str, Any]] = []

    for row_number, values in enumerate(
        reader,
        start=2,
    ):
        record = make_record(
            headers=headers,
            values=list(values),
            row_number=row_number,
        )

        if record:
            records.append(record)

    return records


# ============================================================
# انتخاب روش خواندن فایل
# ============================================================

def read_uploaded_file(
    filename: str,
    file_bytes: bytes,
) -> Tuple[List[Dict[str, Any]], str]:
    extension = Path(filename).suffix.lower()

    if extension == ".xlsx":
        return read_excel_file(file_bytes), "xlsx"

    if extension in {".csv", ".tsv", ".txt"}:
        return (
            read_text_file(
                file_bytes=file_bytes,
                extension=extension,
            ),
            extension[1:],
        )

    raise_bad_request(
        "فرمت فایل مجاز نیست. "
        "فرمت‌های مجاز: xlsx، csv، tsv و txt"
    )

    return [], ""


# ============================================================
# آماده‌سازی سرفصل
# ============================================================

def prepare_topic_data(
    record: Dict[str, Any],
    db: Session,
    seen_ids: set[int],
) -> Tuple[CourseTopic | None, bool]:
    row_number = int(record["_row_number"])

    topic_id = convert_to_int(
        value=record.get("id"),
        field_name="id",
        row_number=row_number,
        minimum=1,
    )

    if topic_id in seen_ids:
        raise_bad_request(
            f"شناسه سرفصل {topic_id} در فایل "
            f"بیش از یک بار تکرار شده است."
        )

    seen_ids.add(topic_id)

    course_id = convert_to_int(
        value=record.get("course_id"),
        field_name="course_id",
        row_number=row_number,
        minimum=1,
    )

    # درس را بررسی می‌کنیم، اگر وجود نداشت، آن را ایجاد می‌کنیم
    course = db.get(Course, course_id)
    if course is None:
        # ایجاد درس جدید با عنوان از فایل یا پیش‌فرض
        course_title = str(record.get("course_title") or "").strip()
        if not course_title:
            course_title = f"درس {course_id}"

        course = Course(
            id=course_id,
            title=course_title,
            description=f"درس ایجادشده خودکار با شناسه {course_id}",
            is_active=True,
        )
        db.add(course)
        db.flush()  # برای اینکه course.id در دسترس باشد

    title = str(
        record.get("title") or ""
    ).strip()

    if not title:
        raise_bad_request(
            f"ستون «title» در سطر "
            f"{row_number} خالی است."
        )

    session_minute = convert_to_int(
        value=record.get("session_minute"),
        field_name="session_minute",
        row_number=row_number,
        default=50,
        minimum=1,
    )

    sort_order = convert_to_int(
        value=record.get("sort_order"),
        field_name="sort_order",
        row_number=row_number,
        default=1,
        minimum=1,
    )

    is_active = convert_to_bool(
        value=record.get("is_active"),
        row_number=row_number,
        default=True,
    )

    description_value = clean_value(
        record.get("description")
    )

    description = (
        str(description_value).strip()
        if not is_empty_value(description_value)
        else None
    )

    existing_topic = db.get(
        CourseTopic,
        topic_id,
    )

    if existing_topic is not None:
        if existing_topic.course_id != course_id:
            raise_bad_request(
                f"سرفصل با شناسه {topic_id} قبلاً "
                "برای درس دیگری ثبت شده است."
            )

        return None, True

    topic = CourseTopic(
        id=topic_id,
        course_id=course_id,
        title=title,
        description=description,
        session_minute=session_minute,
        sort_order=sort_order,
        is_active=is_active,
    )

    return topic, False


# ============================================================
# دریافت لیست درس‌ها
# ============================================================

@router.get(
    "/",
    response_model=List[CourseResponse],
    status_code=status.HTTP_200_OK,
)
def get_courses(
    db: Session = Depends(get_db),
):
    """
    دریافت لیست تمام درس‌های فعال.
    """
    courses = (
        db.query(Course)
        .filter(Course.is_active.is_(True))
        .order_by(Course.id.asc())
        .all()
    )
    return courses


# ============================================================
# ایجاد درس جدید
# ============================================================

@router.post(
    "/",
    response_model=CourseResponse,
    status_code=status.HTTP_201_CREATED,
)
def create_course(
    request: CourseCreate,
    db: Session = Depends(get_db),
):
    """
    ایجاد یک درس جدید با شناسه مشخص.
    """
    existing_course = db.get(Course, request.id)
    if existing_course is not None:
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail=f"درسی با شناسه {request.id} از قبل وجود دارد.",
        )

    course = Course(
        id=request.id,
        title=request.title,
        description=request.description,
        is_active=request.is_active,
    )

    db.add(course)
    db.commit()
    db.refresh(course)

    return course


# ============================================================
# وارد کردن سرفصل‌ها (با ایجاد خودکار درس در صورت نبود)
# ============================================================

@router.post(
    "/import-topics",
    response_model=dict,
)
@router.post(
    "/import_topics",
    response_model=dict,
    include_in_schema=False,
)
async def import_topics(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
):
    if not file.filename:
        raise_bad_request(
            "فایلی انتخاب نشده است."
        )

    filename = Path(file.filename).name
    extension = Path(filename).suffix.lower()

    if extension not in ALLOWED_EXTENSIONS:
        raise_bad_request(
            "فرمت فایل مجاز نیست. "
            "فرمت‌های مجاز: xlsx، csv، tsv و txt"
        )

    file_bytes = await file.read()

    if not file_bytes:
        raise_bad_request(
            "فایل خالی است."
        )

    if len(file_bytes) > MAX_FILE_SIZE:
        raise HTTPException(
            status_code=(
                status.HTTP_413_REQUEST_ENTITY_TOO_LARGE
            ),
            detail="حجم فایل بیشتر از ۱۰ مگابایت است.",
        )

    records, file_type = read_uploaded_file(
        filename=filename,
        file_bytes=file_bytes,
    )

    if not records:
        raise_bad_request(
            "هیچ سطر معتبری در فایل پیدا نشد."
        )

    inserted = 0
    skipped = 0
    seen_ids: set[int] = set()
    topics_to_insert: List[CourseTopic] = []

    try:
        for record in records:
            topic, is_skipped = prepare_topic_data(
                record=record,
                db=db,
                seen_ids=seen_ids,
            )

            if is_skipped:
                skipped += 1
                continue

            if topic is not None:
                topics_to_insert.append(topic)

        if topics_to_insert:
            db.add_all(topics_to_insert)
            db.commit()
            inserted = len(topics_to_insert)

        else:
            db.rollback()

    except HTTPException:
        db.rollback()
        raise

    except IntegrityError as exc:
        db.rollback()

        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail=(
                "یکی از سرفصل‌ها قبلاً ثبت شده یا "
                "ارتباط آن با درس معتبر نیست."
            ),
        ) from exc

    except Exception as exc:
        db.rollback()

        raise HTTPException(
            status_code=(
                status.HTTP_500_INTERNAL_SERVER_ERROR
            ),
            detail=f"خطا در ذخیره سرفصل‌ها: {exc}",
        ) from exc

    return {
        "message": (
            "عملیات وارد کردن سرفصل‌ها "
            "با موفقیت انجام شد."
        ),
        "filename": filename,
        "file_type": file_type,
        "inserted": inserted,
        "skipped": skipped,
        "total_rows": len(records),
    }


# ============================================================
# دریافت سرفصل‌های یک درس (در صورت نبود درس، آرایه خالی)
# ============================================================

@router.get(
    "/{course_id}/topics",
    response_model=List[TopicResponse],
)
def get_topics(
    course_id: int,
    db: Session = Depends(get_db),
):
    if course_id <= 0:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="شناسه درس باید بزرگ‌تر از صفر باشد.",
        )

    # اگر درس وجود نداشت، آرایه خالی برمی‌گردانیم
    course = db.get(Course, course_id)
    if course is None:
        return []  # 200 با آرایه خالی

    topics = (
        db.query(CourseTopic)
        .filter(
            CourseTopic.course_id == course_id,
            CourseTopic.is_active.is_(True),
        )
        .order_by(
            CourseTopic.sort_order.asc(),
            CourseTopic.id.asc(),
        )
        .all()
    )

    return topics